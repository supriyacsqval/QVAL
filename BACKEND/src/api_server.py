import json
import os
import tempfile
from datetime import datetime, timezone
from io import BytesIO

import joblib
import numpy as np
import pandas as pd
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors
from docx import Document as DocxDocument
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from quality_store import build_trend_summary, get_batch_trend, init_store, store_prediction, _connect, _to_float


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.abspath(os.path.join(BASE_DIR, '..'))
ARTIFACTS_DIR = os.path.join(PROJECT_DIR, 'artifacts')
SID_DIR = os.path.join(PROJECT_DIR, 'models', 'latest')

from hybrid_core import (  # type: ignore
    CHARACTERISTICS,
    build_batch_table,
    build_meta_features,
    duration_minutes,
    hybrid_decision,
    load_hybrid_config,
    load_latest_artifacts,
    predict_capa_for_row,
    predict_latest_for_row,
    predict_rf_for_row,
    predict_shap_for_row,
)


app = Flask(__name__)
CORS(
    app,
    resources={
        r"/api/*": {
            "origins": [
                "http://localhost:5173",
                "http://127.0.0.1:5173",
            ]
        }
    },
)


def _load_runtime_artifacts():
    latest_artifacts = load_latest_artifacts(SID_DIR)
    rf_pipeline = joblib.load(os.path.join(ARTIFACTS_DIR, 'rf_batch_pipeline.pkl'))
    cfg = load_hybrid_config(os.path.join(ARTIFACTS_DIR, 'hybrid_config.json'))

    meta_model = None
    meta_path = os.path.join(ARTIFACTS_DIR, 'stacking_meta_model.pkl')
    if os.path.exists(meta_path):
        meta_model = joblib.load(meta_path)

    shap_explainer = None
    try:
        import shap  # type: ignore

        shap_explainer = shap.TreeExplainer(rf_pipeline.named_steps['rf'])
    except Exception:
        shap_explainer = None

    return latest_artifacts, rf_pipeline, cfg, meta_model, shap_explainer


LATEST_ARTIFACTS, RF_PIPELINE, HYBRID_CFG, META_MODEL, SHAP_EXPLAINER = _load_runtime_artifacts()
RF_FEATURE_COLUMNS = HYBRID_CFG['rf_feature_columns']
DB_READY = False

try:
    init_store()
    DB_READY = True
except Exception as exc:
    print(f'Unable to initialize quality store: {exc}')


def _normalize_uploaded_rows(df: pd.DataFrame) -> pd.DataFrame:
    normalized = df.copy()
    aliases = {
        'Batch': ['Batch', 'batch', 'Batch Number', 'BatchNumber', 'UniqueID', 'Unique Id'],
        'ProductDescription': ['ProductDescription', 'Product Description', 'product', 'Product', 'ProductName', 'Product Name', 'Material'],
        'CharacteristicDesc': ['CharacteristicDesc', 'Characteristic', 'TestCharacteristic', 'Test Characteristic', 'Characteristic Name'],
        'Quantative': ['Quantative', 'Quantitative', 'Value', 'value', 'Measure', 'Measurement'],
        'Valuation': ['Valuation', 'Status', 'status', 'Result', 'result'],
        'Qualitative': ['Qualitative', 'qualitative'],
        'StartDate': ['StartDate', 'Start Date', 'start_date', 'Date'],
        'StartTime': ['StartTime', 'Start Time', 'start_time'],
        'EndDate': ['EndDate', 'End Date', 'end_date'],
        'EndTime': ['EndTime', 'End Time', 'end_time'],
        'NotificationID': ['NotificationID', 'Notification Id', 'Notification'],
        'TaskCharacteristic': ['TaskCharacteristic', 'Task Characteristic'],
        'Increase/Decrease': ['Increase/Decrease', 'IncreaseDecrease', 'Direction'],
        'CorrectionValue': ['CorrectionValue', 'Correction Value', 'Correction'],
        'ItemText': ['ItemText', 'Item Text', 'Item'],
        'TaskText': ['TaskText', 'Task Text', 'CAPA', 'CAPA Text', 'ActionText'],
    }

    for target, options in aliases.items():
        if target in normalized.columns:
            continue
        for option in options:
            if option in normalized.columns:
                normalized[target] = normalized[option]
                break

    if 'Batch' not in normalized.columns:
        normalized['Batch'] = 'UNKNOWN_BATCH'
    if 'ProductDescription' not in normalized.columns:
        normalized['ProductDescription'] = 'Unknown Product'
    if 'CharacteristicDesc' not in normalized.columns:
        normalized['CharacteristicDesc'] = 'Unknown Characteristic'
    if 'Valuation' not in normalized.columns:
        normalized['Valuation'] = 'A'

    return normalized


def _build_measure_text(capa_res: dict) -> str:
    characteristic = capa_res.get('capa_characteristic')
    value = capa_res.get('measure_value')
    min_val = capa_res.get('measure_min')
    max_val = capa_res.get('measure_max')
    deviation = capa_res.get('measure_deviation')
    out_of_bounds = bool(capa_res.get('measure_out_of_bounds', False))

    if not characteristic:
        return 'No dominant rejected measure identified.'

    status = 'outside limits' if out_of_bounds else 'near limit'
    return (
        f'{characteristic}: value={value:.2f}, limits=[{min_val:.2f}, {max_val:.2f}], '
        f'deviation={deviation:+.2f} ({status})'
    )


def _build_analysis_text(shap_res: dict, hybrid_res: dict) -> str:
    if shap_res.get('shap_error'):
        return f"Hybrid reason={hybrid_res['reason']}; SHAP unavailable ({shap_res['shap_error']})."

    top_feature = shap_res.get('shap_top_feature')
    top_contrib = shap_res.get('shap_top_contribution')
    top_value = shap_res.get('shap_top_feature_value')
    summary = shap_res.get('shap_summary') or ''
    if top_feature is None:
        return f"Hybrid reason={hybrid_res['reason']}; no SHAP feature contribution available."

    return (
        f"Hybrid reason={hybrid_res['reason']}; top SHAP feature {top_feature} "
        f"(value={top_value:.2f}, contribution={top_contrib:+.4f}). Top factors: {summary}"
    )


def _build_rca_lists(feature_signals: list) -> tuple:
    root_causes = []
    suggestions = []
    warnings = []

    for signal in feature_signals:
        feature = str(signal.get('feature') or '').strip()
        impact = signal.get('impact')
        value = signal.get('value')
        min_val = signal.get('min_val')
        max_val = signal.get('max_val')

        if not feature:
            continue

        try:
            impact = float(impact)
            value = float(value)
            min_val = float(min_val)
            max_val = float(max_val)
        except Exception:
            continue

        if max_val <= min_val:
            continue

        mid = (min_val + max_val) / 2.0

        # Root cause from positive-impact feature contribution.
        if impact > 0:
            root_causes.append(f"{feature}: observed {round(value, 2)}")

            # Suggestion based on where the feature sits relative to limits.
            if value < min_val:
                suggestions.append(f"{feature}: increase towards {round(mid, 2)} (ideal {min_val}-{max_val})")
            elif value > max_val:
                suggestions.append(f"{feature}: decrease towards {round(mid, 2)} (ideal {min_val}-{max_val})")

        # Warning when value is inside but near boundary (10% band).
        if min_val <= value <= max_val:
            band = 0.1 * (max_val - min_val)
            if abs(value - min_val) < band or abs(value - max_val) < band:
                warnings.append(f"{feature}: near boundary (ideal {min_val}-{max_val})")

    # Keep list order stable but remove duplicates.
    root_causes = list(dict.fromkeys(root_causes))
    suggestions = list(dict.fromkeys(suggestions))
    warnings = list(dict.fromkeys(warnings))
    return root_causes, suggestions, warnings


def _build_batch_prediction_result(row, meta_model=None, shap_explainer=None):
    latest_res = predict_latest_for_row(row, LATEST_ARTIFACTS)
    rf_res = predict_rf_for_row(row, RF_PIPELINE, RF_FEATURE_COLUMNS)

    meta_prob = None
    use_meta = bool(float(HYBRID_CFG.get('decision_params', {}).get('use_meta_stacking', 0.0))) and meta_model is not None
    if use_meta:
        meta_input = pd.DataFrame([build_meta_features(latest_res, rf_res)])
        meta_prob = float(meta_model.predict_proba(meta_input)[0][1])

    decision_params = HYBRID_CFG.get('decision_params', {})
    rf_threshold = float(HYBRID_CFG['rf_threshold'])
    meta_threshold = float(decision_params.get('meta_threshold', 0.5))
    hybrid_res = hybrid_decision(
        latest_res,
        rf_res,
        rf_threshold,
        decision_params,
        meta_reject_prob=meta_prob,
        meta_threshold=meta_threshold,
        use_meta_stacking=use_meta,
    )

    shap_res = {
        'shap_top_feature': None,
        'shap_top_contribution': None,
        'shap_top_feature_value': None,
        'shap_summary': None,
        'shap_ranked_features': [],
        'shap_error': '',
    }
    if hybrid_res['status'] == 'R' and shap_explainer is not None:
        shap_res = predict_shap_for_row(row, RF_PIPELINE, RF_FEATURE_COLUMNS, shap_explainer=shap_explainer)

    capa_res = {
        'capa_action': None,
        'capa_characteristic': None,
        'capa_confidence': None,
        'capa_error': '',
        'measure_value': None,
        'measure_min': None,
        'measure_max': None,
        'measure_deviation': None,
        'measure_out_of_bounds': False,
    }
    if hybrid_res['status'] == 'R':
        capa_res = predict_capa_for_row(row, LATEST_ARTIFACTS)

    feature_signals = []
    product_name = row.get('ProductName')

    ranked = shap_res.get('shap_ranked_features') if isinstance(shap_res, dict) else None
    if isinstance(ranked, list):
        for item in ranked:
            feature_name = item.get('feature')
            if not feature_name:
                continue
            bounds = LATEST_ARTIFACTS['reference_lookup'].get((product_name, feature_name))
            if not bounds:
                continue
            signal_val = row.get(feature_name)
            if pd.isna(signal_val):
                continue
            feature_signals.append(
                {
                    'feature': feature_name,
                    'impact': abs(float(item.get('impact', 0.0))),
                    'value': float(signal_val),
                    'min_val': float(bounds[0]),
                    'max_val': float(bounds[1]),
                }
            )

    if not feature_signals and capa_res.get('capa_characteristic'):
        feature_signals.append(
            {
                'feature': capa_res.get('capa_characteristic'),
                'impact': 1.0,
                'value': capa_res.get('measure_value'),
                'min_val': capa_res.get('measure_min'),
                'max_val': capa_res.get('measure_max'),
            }
        )

    root_causes, suggestions, warnings = _build_rca_lists(feature_signals)

    source_capa_details = {
        'notification_id': row.get('NotificationID'),
        'task_characteristic': row.get('TaskCharacteristic'),
        'direction': row.get('Increase/Decrease'),
        'correction_value': row.get('CorrectionValue'),
        'item_text': row.get('ItemText'),
        'task_text': row.get('TaskText'),
    }

    return {
        'batch': row.get('UniqueID') or row.get('Batch'),
        'product': row.get('ProductName'),
        'start_time': row.get('StartTime'),
        'end_time': row.get('EndTime'),
        'latest': latest_res,
        'rf': rf_res,
        'hybrid': hybrid_res,
        'meta_reject_probability': round(float(meta_prob), 4) if meta_prob is not None else None,
        'decision_text': (
            f"Batch {row.get('UniqueID') or row.get('Batch')} predicted {hybrid_res['status']} "
            f"(latest={latest_res['status']}:{float(latest_res['reject_confidence']):.3f}, "
            f"rf={rf_res['status']}:{float(rf_res['reject_confidence']):.3f})."
        ),
        'measure_text': 'Not rejected; no CAPA measure required.' if hybrid_res['status'] != 'R' else _build_measure_text(capa_res),
        'analysis_text': _build_analysis_text(shap_res, hybrid_res),
        'root_causes': root_causes,
        'suggestions': suggestions,
        'warnings': warnings,
        'source_capa_details': source_capa_details,
        'source_capa_text': (row.get('TaskText') or row.get('ItemText') or ''),
        'shap': {
            'enabled': hybrid_res['status'] == 'R' and shap_explainer is not None,
            'top_feature': shap_res.get('shap_top_feature'),
            'top_contribution': shap_res.get('shap_top_contribution'),
            'top_feature_value': shap_res.get('shap_top_feature_value'),
            'summary': shap_res.get('shap_summary'),
            'error': shap_res.get('shap_error') or '',
        },
        'capa': {
            'action': capa_res.get('capa_action'),
            'characteristic': capa_res.get('capa_characteristic'),
            'confidence': capa_res.get('capa_confidence'),
            'error': capa_res.get('capa_error') or '',
            'measure_value': capa_res.get('measure_value'),
            'measure_min': capa_res.get('measure_min'),
            'measure_max': capa_res.get('measure_max'),
            'measure_deviation': capa_res.get('measure_deviation'),
            'measure_out_of_bounds': bool(capa_res.get('measure_out_of_bounds', False)),
        },
    }


@app.route('/api/health', methods=['GET'])
def health():
    return jsonify(
        {
            'ok': True,
            'hybrid_ready': True,
            'store_ready': DB_READY,
            'rf_artifact': os.path.exists(os.path.join(ARTIFACTS_DIR, 'rf_batch_pipeline.pkl')),
            'meta_artifact': os.path.exists(os.path.join(ARTIFACTS_DIR, 'stacking_meta_model.pkl')),
            'shap_ready': SHAP_EXPLAINER is not None,
        }
    )


@app.route('/api/options', methods=['GET'])
def options():
    reference_lookup = LATEST_ARTIFACTS['reference_lookup']
    products_dict = {}
    for prod, char in reference_lookup.keys():
        products_dict.setdefault(prod, [])
        if char not in products_dict[prod]:
            products_dict[prod].append(char)
    return jsonify(products_dict)


@app.route('/api/options/detail', methods=['GET'])
def options_detail():
    reference_lookup = LATEST_ARTIFACTS['reference_lookup']
    products: dict = {}

    for (prod, char), bounds in reference_lookup.items():
        min_val, max_val = bounds
        products.setdefault(prod, [])
        products[prod].append(
            {
                'characteristic': char,
                'min_value': float(min_val),
                'max_value': float(max_val),
                'target': float((min_val + max_val) / 2.0),
            }
        )

    product_names = sorted(products.keys())
    characteristic_names = sorted({item['characteristic'] for values in products.values() for item in values})

    return jsonify(
        {
            'products': products,
            'products_list': product_names,
            'characteristics_list': characteristic_names,
        }
    )


@app.route('/api/suggestions', methods=['GET'])
def suggestions():
    kind = (request.args.get('type') or '').strip().lower()
    query = (request.args.get('q') or '').strip().lower()
    product = (request.args.get('product') or '').strip()
    limit = request.args.get('limit', default=12, type=int)
    limit = max(1, min(limit, 100))

    if kind not in {'product', 'characteristic', 'batch'}:
        return jsonify({'error': 'type must be product, characteristic, or batch'}), 400

    items = []

    if kind in {'product', 'characteristic'}:
        reference_lookup = LATEST_ARTIFACTS['reference_lookup']
        if kind == 'product':
            items = sorted({prod for prod, _ in reference_lookup.keys()})
        else:
            if product:
                items = sorted({char for prod, char in reference_lookup.keys() if prod == product})
            else:
                items = sorted({char for _, char in reference_lookup.keys()})
    else:
        if not DB_READY:
            return jsonify({'items': []})
        try:
            import sqlite3
            from quality_store import DB_PATH  # type: ignore

            with sqlite3.connect(DB_PATH) as conn:
                rows = conn.execute(
                    '''
                    SELECT DISTINCT batch
                    FROM quality_history
                    WHERE batch IS NOT NULL AND TRIM(batch) != ''
                    ORDER BY batch
                    LIMIT 500
                    '''
                ).fetchall()
                items = [str(row[0]) for row in rows]
        except Exception:
            items = []

    if query:
        items = [item for item in items if query in item.lower()]

    return jsonify({'items': items[:limit]})


@app.route('/api/trend', methods=['GET'])
def trend():
    product = (request.args.get('product') or '').strip()
    characteristic = (request.args.get('characteristic') or '').strip() or None
    limit = request.args.get('limit', default=8, type=int)

    if not DB_READY:
        return jsonify({'error': 'Trend store is not available.'}), 503

    if not product:
        # Return summary for all products
        with _connect() as conn:
            df = pd.read_sql_query(
                '''
                SELECT
                    product,
                    COUNT(*) AS total_rows,
                    SUM(CASE WHEN status = 'R' THEN 1 ELSE 0 END) AS reject_rows,
                    AVG(quantitative) AS avg_quantitative,
                    AVG(deviation) AS avg_deviation
                FROM quality_history
                WHERE product IS NOT NULL AND product != ''
                GROUP BY product
                ORDER BY total_rows DESC
                ''',
                conn,
            )
        
        products_summary = []
        for _, row in df.iterrows():
            total = int(row['total_rows'] or 0)
            reject = int(row['reject_rows'] or 0)
            products_summary.append({
                'product': row['product'],
                'total_rows': total,
                'reject_rows': reject,
                'reject_rate': (reject / total) if total else 0.0,
                'avg_quantitative': _to_float(row['avg_quantitative']),
                'avg_deviation': _to_float(row['avg_deviation']),
            })
        
        return jsonify({
            'all_products': True,
            'products': products_summary,
            'total_products': len(products_summary),
            'trend_text': f'Found {len(products_summary)} products in the database.',
        })

    return jsonify(build_trend_summary(product, characteristic=characteristic, limit=limit))


@app.route('/api/predict', methods=['POST'])
def predict():
    payload = request.get_json(silent=True) or {}
    product = (payload.get('product') or '').strip()
    characteristics = payload.get('characteristics', {})
    start_time_str = (payload.get('startTime') or '').strip()
    end_time_str = (payload.get('endTime') or '').strip()
    batch = (payload.get('batch') or '').strip()

    if not all([product, characteristics, start_time_str, end_time_str]):
        return jsonify({'error': 'Missing input fields.'}), 400

    if not isinstance(characteristics, dict):
        return jsonify({'error': 'characteristics must be a dict.'}), 400

    # Create row
    row = {
        'ProductName': product,
        'StartTime': start_time_str,
        'EndTime': end_time_str,
        'Batch': batch,
    }
    for char in CHARACTERISTICS:
        val = characteristics.get(char)
        if val is not None:
            try:
                row[char] = float(val)
            except Exception:
                return jsonify({'error': f'Invalid value for {char}.'}), 400
        else:
            row[char] = np.nan

    # Predict
    latest_result = predict_latest_for_row(pd.Series(row), LATEST_ARTIFACTS)
    status_label = latest_result['status']
    decision_reason = latest_result.get('decision_reason', 'model')

    duration_min = duration_minutes(start_time_str, end_time_str)

    # For derived, use the first characteristic with value
    first_char = None
    quantitative = None
    min_val = None
    max_val = None
    for char in CHARACTERISTICS:
        if not pd.isna(row[char]):
            first_char = char
            quantitative = row[char]
            bounds = LATEST_ARTIFACTS['reference_lookup'].get((product, char))
            if bounds:
                min_val, max_val = bounds
            break

    if first_char:
        target = (min_val + max_val) / 2.0 if min_val is not None and max_val is not None else quantitative
        deviation = quantitative - target
        is_out_of_bounds = quantitative < min_val or quantitative > max_val if min_val is not None and max_val is not None else False
    else:
        deviation = 0
        is_out_of_bounds = False

    # Trend summary for the first characteristic
    trend_summary = build_trend_summary(product, first_char, limit=8) if DB_READY and first_char else {}

    # CAPA
    capa_action = None
    if status_label == 'R' and first_char and not is_out_of_bounds:
        capa_input = pd.DataFrame([
            {
                'ProductDescription': product,
                'CharacteristicDesc': first_char,
                'Quantative': quantitative,
                'MinValue': min_val,
                'MaxValue': max_val,
                'Duration_min': duration_min,
                'DeviationValue': deviation,
            }
        ])
        x_capa = LATEST_ARTIFACTS['preprocessor_capa'].transform(capa_input)
        pred_capa_probs = LATEST_ARTIFACTS['model_capa'].predict(x_capa)[0]
        pred_capa_idx = int(pd.Series(pred_capa_probs).idxmax())
        capa_action = LATEST_ARTIFACTS['le_capa'].inverse_transform([pred_capa_idx])[0]

    decision_text = (
        f'Batch for {product} was {"rejected" if status_label == "R" else "accepted"} based on all characteristics.'
    )
    measure_text = (
        f'All characteristics evaluated for {product}.'
    )
    analysis_text = (
        f'RCA focus: Batch evaluation for {product}. Trend context: {trend_summary.get("trend_text", "")}'
    )

    manual_signals = []
    for char in CHARACTERISTICS:
        if not pd.isna(row[char]):
            bounds = LATEST_ARTIFACTS['reference_lookup'].get((product, char))
            if bounds:
                min_v, max_v = bounds
                impact = abs(float(row[char] - (min_v + max_v)/2)) / max(1e-6, (max_v - min_v))
                manual_signals.append({
                    'feature': char,
                    'impact': impact,
                    'value': row[char],
                    'min_val': min_v,
                    'max_val': max_v,
                })

    root_causes, suggestions, warnings = _build_rca_lists(manual_signals)

    rca_parts = [
        f'RCA Report for {product}.',
        decision_text,
        measure_text,
        analysis_text,
    ]
    if root_causes:
        rca_parts.append('Root causes: ' + '; '.join(root_causes) + '.')
    if suggestions:
        rca_parts.append('Suggestions: ' + '; '.join(suggestions) + '.')
    if warnings:
        rca_parts.append('Warnings: ' + '; '.join(warnings) + '.')
    rca_report = ' '.join(rca_parts)

    response = {
        'status': status_label,
        'decision_reason': decision_reason,
        'decision_text': decision_text,
        'measure_text': measure_text,
        'analysis_text': analysis_text,
        'rca_report': rca_report,
        'root_causes': root_causes,
        'suggestions': suggestions,
        'warnings': warnings,
        'trend_text': trend_summary.get('trend_text', ''),
        'capa_action': capa_action,
        'capa_text': capa_action or 'No CAPA suggested',
        'trend': trend_summary,
        'derived': {
            'duration_min': duration_min,
            'min_val': min_val or 0,
            'max_val': max_val or 0,
            'deviation': deviation,
        },
        'shap': {
            'enabled': False,
            'error': 'SHAP not available for manual input with multiple characteristics',
        },
        'capa': {
            'action': capa_action,
            'characteristic': first_char,
        } if capa_action else None,
    }

    if DB_READY:
        try:
            store_prediction(
                {
                    'batch': batch,
                    'product': product,
                    'characteristic': first_char or 'All',
                    'start_time': start_time_str,
                    'end_time': end_time_str,
                    'quantitative': quantitative,
                    'min_value': min_val,
                    'max_value': max_val,
                    'duration_min': duration_min,
                    'deviation': deviation,
                    'status': status_label,
                    'decision_reason': decision_reason,
                    'decision_text': decision_text,
                    'measure_text': measure_text,
                    'analysis_text': analysis_text,
                    'capa_action': capa_action,
                    'capa_confidence': None,
                    'trend_text': trend_summary.get('trend_text') if trend_summary else None,
                    'record_source': 'manual_prediction_all',
                    'event_time': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                    'payload_json': response,
                }
            )
        except Exception as store_exc:
            response['store_error'] = str(store_exc)

    return jsonify(response)


@app.route('/api/batch/predict', methods=['POST'])
def batch_predict():
    if not os.path.exists(os.path.join(ARTIFACTS_DIR, 'rf_batch_pipeline.pkl')):
        return jsonify({'error': 'RF artifact not available.'}), 503

    uploaded_df = None
    source_name = None

    if 'file' in request.files and request.files['file'].filename:
        file_obj = request.files['file']
        source_name = file_obj.filename
        try:
            uploaded_df = pd.read_csv(file_obj)
        except Exception as exc:
            return jsonify({'error': f'Unable to read CSV file: {exc}'}), 400
    else:
        payload = request.get_json(silent=True) or {}
        rows = payload.get('rows')
        if isinstance(rows, list) and rows:
            uploaded_df = pd.DataFrame(rows)
            source_name = payload.get('name')

    if uploaded_df is None or uploaded_df.empty:
        return jsonify({'error': 'Upload a CSV file or send a JSON body with rows.'}), 400

    normalized_df = _normalize_uploaded_rows(uploaded_df)
    temp_path = None

    try:
        with tempfile.NamedTemporaryFile('w', suffix='.csv', delete=False, encoding='utf-8', newline='') as temp_file:
            temp_path = temp_file.name
            normalized_df.to_csv(temp_file, index=False)

        batch_df = build_batch_table(temp_path, None)

        results = []
        for _, row in batch_df.iterrows():
            result = _build_batch_prediction_result(row, meta_model=META_MODEL, shap_explainer=SHAP_EXPLAINER)
            results.append(result)

            if DB_READY:
                try:
                    source_capa = result.get('source_capa_details') or {}
                    store_prediction(
                        {
                            'batch': result['batch'],
                            'product': result['product'],
                            'characteristic': None,
                            'start_time': result['start_time'],
                            'end_time': result['end_time'],
                            'quantitative': None,
                            'min_value': None,
                            'max_value': None,
                            'duration_min': None,
                            'deviation': None,
                            'status': result['hybrid']['status'],
                            'decision_reason': result['hybrid']['reason'],
                            'decision_text': result['decision_text'],
                            'measure_text': result['measure_text'],
                            'analysis_text': result['analysis_text'],
                            'capa_action': result['capa']['action'],
                            'capa_confidence': result['capa']['confidence'],
                            'trend_text': result['analysis_text'],
                            'record_source': 'batch_prediction',
                            'event_time': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                            'payload_json': {
                                **result,
                                'source_capa_details': source_capa,
                            },
                        }
                    )
                except Exception as store_exc:
                    result['store_error'] = str(store_exc)

        accepted = sum(1 for item in results if item['hybrid']['status'] == 'A')
        rejected = sum(1 for item in results if item['hybrid']['status'] == 'R')

        for item in results:
            capa_text = f"Recommended CAPA: {item['capa']['action']}" if item['capa'].get('action') else 'No CAPA action required.'
            item_root = item.get('root_causes') or []
            item_suggestions = item.get('suggestions') or []
            item_warnings = item.get('warnings') or []
            source_capa_text = (item.get('source_capa_text') or '').strip()
            item_parts = [
                item.get('decision_text', ''),
                item.get('measure_text', ''),
                item.get('analysis_text', ''),
                capa_text,
            ]
            if source_capa_text:
                item_parts.append('Historical CAPA detail: ' + source_capa_text)
            if item_root:
                item_parts.append('Root causes: ' + '; '.join(item_root) + '.')
            if item_suggestions:
                item_parts.append('Suggestions: ' + '; '.join(item_suggestions) + '.')
            if item_warnings:
                item_parts.append('Warnings: ' + '; '.join(item_warnings) + '.')
            item['rca_report'] = ' '.join(part for part in item_parts if part).strip()

        top_rejects = [item for item in results if item['hybrid']['status'] == 'R'][:5]
        if top_rejects:
            highlights = '; '.join(
                f"{item.get('batch')}: {item.get('hybrid', {}).get('reason', 'model_reject')}"
                for item in top_rejects
            )
            batch_rca_report = (
                f'Batch RCA Summary: {rejected} rejected of {len(results)} batches. '
                f'Top rejection signals -> {highlights}.'
            )
        else:
            batch_rca_report = f'Batch RCA Summary: all {len(results)} batches accepted; no reject root cause identified.'

        return jsonify(
            {
                'source_file': source_name,
                'rows_received': int(len(uploaded_df)),
                'batches': int(len(results)),
                'accepted': accepted,
                'rejected': rejected,
                'rca_report': batch_rca_report,
                'results': results,
            }
        )
    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass


def _generate_pdf_report(analysis_data: dict) -> bytes:
    """Generate a PDF report from analysis data."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#1a1a1a'),
        spaceAfter=30,
        alignment=1
    )
    
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=14,
        textColor=colors.HexColor('#2c3e50'),
        spaceAfter=12,
        spaceBefore=12
    )
    
    body_style = ParagraphStyle(
        'CustomBody',
        parent=styles['BodyText'],
        fontSize=11,
        spaceAfter=10
    )
    
    # Title
    elements.append(Paragraph("Quality Analysis Report", title_style))
    elements.append(Spacer(1, 0.2*inch))
    
    # Header Info
    if 'decision_text' in analysis_data:
        elements.append(Paragraph("Decision & Analysis", heading_style))
        elements.append(Paragraph(analysis_data['decision_text'], body_style))
        elements.append(Spacer(1, 0.15*inch))
    
    # Measure Details
    if 'measure_text' in analysis_data:
        elements.append(Paragraph("Measurement Details", heading_style))
        elements.append(Paragraph(analysis_data['measure_text'], body_style))
        elements.append(Spacer(1, 0.15*inch))
    
    # Analysis Section
    if 'analysis_text' in analysis_data:
        elements.append(Paragraph("SHAP Analysis", heading_style))
        elements.append(Paragraph(analysis_data['analysis_text'], body_style))
        elements.append(Spacer(1, 0.15*inch))
    
    # CAPA Section
    if 'capa_text' in analysis_data:
        elements.append(Paragraph("CAPA Action", heading_style))
        elements.append(Paragraph(analysis_data['capa_text'], body_style))
        elements.append(Spacer(1, 0.15*inch))
    
    # Trend Analysis
    if 'trend_text' in analysis_data and analysis_data['trend_text']:
        elements.append(Paragraph("Trend Analysis", heading_style))
        elements.append(Paragraph(analysis_data['trend_text'], body_style))
        elements.append(Spacer(1, 0.15*inch))
    
    # Metadata
    elements.append(Spacer(1, 0.2*inch))
    elements.append(Paragraph("Report Generated", heading_style))
    elements.append(Paragraph(f"Date: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}", body_style))
    
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


def _generate_docx_report(analysis_data: dict) -> bytes:
    """Generate a DOCX report from analysis data."""
    doc = DocxDocument()
    
    # Add title
    title = doc.add_heading('Quality Analysis Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add decision section
    if 'decision_text' in analysis_data:
        doc.add_heading('Decision & Analysis', level=1)
        doc.add_paragraph(analysis_data['decision_text'])
        doc.add_paragraph()
    
    # Add measurement details
    if 'measure_text' in analysis_data:
        doc.add_heading('Measurement Details', level=1)
        doc.add_paragraph(analysis_data['measure_text'])
        doc.add_paragraph()
    
    # Add SHAP analysis
    if 'analysis_text' in analysis_data:
        doc.add_heading('SHAP Analysis', level=1)
        doc.add_paragraph(analysis_data['analysis_text'])
        doc.add_paragraph()
    
    # Add CAPA action
    if 'capa_text' in analysis_data:
        doc.add_heading('CAPA Action', level=1)
        doc.add_paragraph(analysis_data['capa_text'])
        doc.add_paragraph()
    
    # Add trend analysis
    if 'trend_text' in analysis_data and analysis_data['trend_text']:
        doc.add_heading('Trend Analysis', level=1)
        doc.add_paragraph(analysis_data['trend_text'])
        doc.add_paragraph()
    
    # Add metadata
    doc.add_heading('Report Metadata', level=1)
    doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}")
    
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


@app.route('/api/export', methods=['POST'])
def export_report():
    """Export analysis results as PDF or DOCX."""
    payload = request.get_json(silent=True) or {}
    analysis_data = payload.get('analysis', {})
    export_format = (payload.get('format', 'pdf') or 'pdf').lower().strip()
    
    if not analysis_data:
        return jsonify({'error': 'Missing analysis data.'}), 400
    
    if export_format not in ['pdf', 'docx']:
        return jsonify({'error': 'Format must be "pdf" or "docx".'}), 400
    
    try:
        if export_format == 'pdf':
            content = _generate_pdf_report(analysis_data)
            mimetype = 'application/pdf'
            filename = f"quality_report_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.pdf"
        else:
            content = _generate_docx_report(analysis_data)
            mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            filename = f"quality_report_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.docx"
        
        return send_file(
            BytesIO(content),
            mimetype=mimetype,
            as_attachment=True,
            download_name=filename
        )
    except Exception as exc:
        return jsonify({'error': f'Export failed: {str(exc)}'}), 500


@app.route('/api/batch/trend', methods=['GET'])
def batch_trend():
    """Get trend data for a specific batch/tablet across time."""
    batch = (request.args.get('batch') or '').strip()
    if not batch:
        return jsonify({'error': 'Missing batch parameter.'}), 400
    if not DB_READY:
        return jsonify({'error': 'Trend store is not available.'}), 503
    
    try:
        result = get_batch_trend(batch, limit=100)
        return jsonify(result)
    except Exception as exc:
        return jsonify({'error': f'Failed to fetch batch trend: {str(exc)}'}), 500


if __name__ == '__main__':
    app.run(debug=True, port=5050)